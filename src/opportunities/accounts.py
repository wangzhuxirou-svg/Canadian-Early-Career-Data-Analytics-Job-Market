import re
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook


_USERNAME_PATTERN = re.compile(r"^[A-Za-z0-9._]{1,30}$")
_REQUIRED_HEADERS = {
    "username": "Instagram 账号名称",
    "name": "社团名称",
    "profile_url": "Instagram 主页链接",
}


def _organizer_id(username: str) -> str:
    organizer_id = re.sub(r"[^a-z0-9]+", "_", username.lower()).strip("_")
    if not organizer_id or not organizer_id[0].isalpha():
        organizer_id = f"club_{organizer_id}"
    return organizer_id


def _cell_text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def import_organizers_xlsx(path: Path) -> list[dict]:
    """Import the account table without treating selection as access verification."""
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        worksheet = next(
            (
                sheet
                for sheet in workbook.worksheets
                if any(
                    _cell_text(cell.value) == _REQUIRED_HEADERS["username"]
                    for row in sheet.iter_rows()
                    for cell in row
                )
            ),
            None,
        )
        if worksheet is None:
            raise ValueError("Workbook is missing the Instagram account header row.")

        header_row_number = next(
            row_number
            for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1)
            if _REQUIRED_HEADERS["username"] in {_cell_text(value) for value in row}
        )
        header_values = next(
            worksheet.iter_rows(
                min_row=header_row_number,
                max_row=header_row_number,
                values_only=True,
            )
        )
        headers = {_cell_text(value): index for index, value in enumerate(header_values)}
        missing_headers = set(_REQUIRED_HEADERS.values()) - set(headers)
        if missing_headers:
            raise ValueError(f"Workbook is missing required columns: {sorted(missing_headers)}.")

        records = []
        seen_usernames = set()
        seen_ids = set()
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=header_row_number + 1, values_only=True),
            start=header_row_number + 1,
        ):
            username = _cell_text(values[headers[_REQUIRED_HEADERS["username"]]])
            if username.startswith("@"):
                username = username[1:]
            name = _cell_text(values[headers[_REQUIRED_HEADERS["name"]]])
            profile_url = _cell_text(values[headers[_REQUIRED_HEADERS["profile_url"]]])
            if not username and not name and not profile_url:
                continue
            if not name and not profile_url and not _USERNAME_PATTERN.fullmatch(username):
                continue
            if not username or not name or not profile_url:
                raise ValueError(f"Workbook row {row_number} has an incomplete account record.")
            if not _USERNAME_PATTERN.fullmatch(username):
                raise ValueError(f"Workbook row {row_number} has an invalid Instagram username.")
            if username.casefold() in seen_usernames:
                raise ValueError(f"Workbook contains duplicate username {username!r}.")

            parsed_url = urlparse(profile_url)
            profile_path = parsed_url.path.strip("/").split("/")
            if (
                parsed_url.scheme != "https"
                or parsed_url.netloc.casefold() not in {"instagram.com", "www.instagram.com"}
                or len(profile_path) != 1
                or profile_path[0].casefold() != username.casefold()
                or parsed_url.query
                or parsed_url.fragment
            ):
                raise ValueError(
                    f"Workbook row {row_number} profile URL does not match username {username!r}."
                )

            organizer_id = _organizer_id(username)
            if organizer_id in seen_ids:
                raise ValueError(f"Workbook creates duplicate organizer ID {organizer_id!r}.")
            seen_usernames.add(username.casefold())
            seen_ids.add(organizer_id)
            records.append({
                "organizer_id": organizer_id,
                "name": name,
                "username": username,
                "profile_url": f"https://www.instagram.com/{username}/",
                "enabled": True,
                "priority": len(records) + 1,
                "last_verified_at": None,
            })

        if not records:
            raise ValueError("Workbook contains no Instagram account records.")
        return records
    finally:
        workbook.close()
